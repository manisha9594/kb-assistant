"""Document ingestion: PDF, DOCX, TXT, and Markdown -> chunks -> embeddings -> ChromaDB."""
import hashlib
import io
from collections import defaultdict
from pathlib import Path

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.config import get_settings
from app.providers import get_vectorstore

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


class IngestionError(ValueError):
    pass


def _load_pdf(data: bytes, name: str) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return [
        Document(page_content=text, metadata={"source": name, "page": i + 1})
        for i, page in enumerate(reader.pages)
        if (text := (page.extract_text() or "").strip())
    ]


def _load_docx(data: bytes, name: str) -> list[Document]:
    import docx

    paragraphs = [p.text for p in docx.Document(io.BytesIO(data)).paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    return [Document(page_content=text, metadata={"source": name, "page": 0})] if text else []


def _load_text(data: bytes, name: str) -> list[Document]:
    text = data.decode("utf-8", errors="ignore").strip()
    return [Document(page_content=text, metadata={"source": name, "page": 0})] if text else []


LOADERS = {".pdf": _load_pdf, ".docx": _load_docx, ".txt": _load_text, ".md": _load_text}


def load_file(data: bytes, filename: str) -> list[Document]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise IngestionError(f"Unsupported file type '{ext}'. Upload PDF, DOCX, TXT, or MD.")
    try:
        docs = LOADERS[ext](data, filename)
    except Exception as exc:
        raise IngestionError(f"Could not read '{filename}'. The file may be corrupted.") from exc
    if not docs:
        raise IngestionError(f"No text found in '{filename}'. Scanned PDFs need OCR first.")
    return docs


def split(docs: list[Document]) -> list[Document]:
    s = get_settings()
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=s.chunk_size,
        chunk_overlap=s.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(docs)
    for i, c in enumerate(chunks):
        c.metadata["chunk"] = i
    return chunks


def _chunk_id(c: Document) -> str:
    key = f"{c.metadata['source']}|{c.metadata['page']}|{c.metadata['chunk']}"
    return hashlib.sha256(key.encode()).hexdigest()[:32]


def delete_document(filename: str) -> int:
    vs = get_vectorstore()
    ids = vs.get(where={"source": filename}).get("ids", [])
    if ids:
        vs.delete(ids=ids)
    return len(ids)


def ingest(data: bytes, filename: str) -> dict:
    """Index a file. Re-uploading the same filename replaces its previous chunks."""
    chunks = split(load_file(data, filename))
    delete_document(filename)
    get_vectorstore().add_documents(chunks, ids=[_chunk_id(c) for c in chunks])
    pages = len({c.metadata["page"] for c in chunks})
    return {"filename": filename, "pages": pages, "chunks": len(chunks)}


def list_documents() -> list[dict]:
    metas = get_vectorstore().get(include=["metadatas"]).get("metadatas", [])
    stats: dict[str, dict] = defaultdict(lambda: {"chunks": 0, "pages": set()})
    for m in metas:
        stats[m["source"]]["chunks"] += 1
        stats[m["source"]]["pages"].add(m["page"])
    return [
        {"filename": n, "pages": len(v["pages"]), "chunks": v["chunks"]}
        for n, v in sorted(stats.items())
    ]


def search(query: str, k: int | None = None) -> list[dict]:
    """Vector search returning chunks above the relevance threshold, best first."""
    s = get_settings()
    vs = get_vectorstore()
    if not vs.get(limit=1).get("ids"):
        return []
    results = vs.similarity_search_with_relevance_scores(query, k=k or s.top_k)
    return [
        {
            "kind": "document",
            "source": d.metadata["source"],
            "page": d.metadata["page"],
            "text": d.page_content,
            "score": round(score, 3),
        }
        for d, score in results
        if score >= s.min_relevance
    ]
